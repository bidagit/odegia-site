# -*- coding: utf-8 -*-
"""Grille d entretien du diagnostic, generee pour un dossier client.

    python scripts/diagnostic/grille.py "<dossier>"

Lit le nom du client et la date dans entree.json et ecrit grille-entretien.docx
dans le dossier. La grille est la meme pour tous les clients, seule l en-tete
change. C est ce qui rend le diagnostic reproductible, une grille reecrite a la
main pour chaque client finit par reinterpreter la methode.

Elle suit le protocole du vault, `Odegia - Protocole de diagnostic`, et porte
sa version. Toute question modifiee ici change la version de la methode.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERT = RGBColor(0x0F, 0x7A, 0x5C)
GRIS = RGBColor(0x5C, 0x64, 0x62)

dossier = Path(sys.argv[1]).resolve()
entree = json.loads((dossier / "entree.json").read_text(encoding="utf-8"))
calcul = (Path(__file__).parent / "calcul.mts").read_text(encoding="utf-8")
VERSION = re.search(r'METHODE_VERSION = "([^"]+)"', calcul).group(1)
client = entree["client"]["nom"]
date = entree.get("dateEntretien") or "date à fixer"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.5)
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10)
pied = sec.footer.paragraphs[0].add_run(f"Odegia · grille de diagnostic · méthode v{VERSION}")
pied.font.size = Pt(7); pied.font.color.rgb = GRIS


def fenetre(t):
    # Largeur 100 % de la fenetre, regle du vault pour toutes les tables Word.
    pr = t._tbl.tblPr
    for old in pr.findall(qn("w:tblW")):
        pr.remove(old)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct")
    pr.append(w)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER


def titre(txt, sous=None):
    r = doc.add_paragraph().add_run(txt); r.bold = True; r.font.size = Pt(15)
    if sous:
        r2 = doc.add_paragraph().add_run(sous); r2.font.size = Pt(9); r2.font.color.rgb = GRIS


def intertitre(txt):
    r = doc.add_paragraph().add_run(txt); r.bold = True


def horloge():
    r = doc.add_paragraph().add_run("Début ____h____     Fin ____h____")
    r.bold = True; r.font.color.rgb = VERT


def table2(lignes, haut=1.0):
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for a, b in lignes:
        c = t.add_row().cells
        c[0].text, c[1].text = a, b
        for r in c[0].paragraphs[0].runs: r.bold = True; r.font.size = Pt(9.5)
        for r in c[1].paragraphs[0].runs: r.font.color.rgb = GRIS; r.font.size = Pt(8.5)
        c[0].width, c[1].width = Cm(6), Cm(11.4)
        t.rows[-1].height = Cm(haut)
    fenetre(t)


def grille(entetes, n, haut=0.8, largeurs=None):
    t = doc.add_table(rows=1, cols=len(entetes)); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, entetes):
        c.text = h
        for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(8.5)
    for _ in range(n):
        t.add_row().height = Cm(haut)
    if largeurs:
        for row in t.rows:
            for c, w in zip(row.cells, largeurs):
                c.width = Cm(w)
    fenetre(t)


# ── Page 1, reperes ─────────────────────────────────────────────────────────
titre(f"Diagnostic {client}, {date}",
      "Demander l'accord pour enregistrer avant de lancer. On note les chiffres et les cases, "
      "l'enregistrement garde le reste. On ne chiffre rien devant le client.")
intertitre("Déroulé")
table2([
    ("Questionnaire", "20 min à l'oral, 5 min s'il a été rempli avant"),
    ("Inventaire et sélection", "15 min, tableau page 3"),
    ("Une fiche par tâche retenue", "12 min au plus chacune, 4 à 5 tâches"),
    ("Fin", "10 min, questions de clôture"),
], 0.8)
intertitre("Les mesures, à faire préciser sur le moment")
table2([
    ("Occurrence", "Une exécution complète, du déclencheur à la sortie remise au destinataire."),
    ("Fois par an", "Nombre d'occurrences sur douze mois, saisons comprises. "
                    "« Combien de fois l'an dernier ? » plutôt que « en général ? »."),
    ("Durée", "Temps de travail actif d'une occurrence ordinaire, relecture et recherche comprises, "
              "attente exclue. Partir de la dernière fois, en minutes."),
    ("Acteur", "La personne qui fait la tâche. Une tâche à deux se relève en deux lignes."),
    ("Bascule", "Temps pour s'y remettre avant et s'en défaire après. 20 % de la durée active, "
                "annoncés au client, qui corrige tâche par tâche entre 0 et 30 %. Compté dans le coût et le gain."),
    ("Taux", "Ce que le client déclare pour une heure de cet acteur. "
             "Sans réponse, 60 € de l'heure, annoncés comme hypothèse dans le rapport."),
], 1.25)

# ── Page 2, questionnaire ───────────────────────────────────────────────────
doc.add_page_break()
titre("Questionnaire", "Partie A de la fiche. Une réponse approximative suffit.")
horloge()
intertitre("Étape 0, définir, 5 minutes")
table2([
    ("Ce que l'organisation produit", "une sortie mesurable, pas une promesse"),
    ("Pour qui, exactement", "la description la plus étroite encore vraie"),
    ("Par quel flux la valeur se fabrique", "la séquence principale, premier candidat au parcours"),
], 1.1)
intertitre("Le reste du questionnaire")
table2([
    ("Activité en une phrase", ""),
    ("Forme juridique", ""),
    ("Personnes, prestataires réguliers compris", ""),
    ("Chiffre d'affaires ou budget annuel", "fourchette"),
    ("TVA", "assujetti / franchise / non concerné"),
    ("Qui fait l'administratif", "noms et fonctions"),
    ("Heures d'admin par semaine, tous confondus", "subies ou choisies"),
    ("Taux d'une heure, par acteur", "déclaré par le client"),
    ("Facturation et devis", "outil"),
    ("Banque", ""),
    ("Messagerie et agenda", ""),
    ("Site et formulaires", ""),
    ("Signature électronique", ""),
    ("Stockage des documents", ""),
    ("Comptabilité", "interne / cabinet"),
    ("Autre outil quotidien", ""),
], 0.85)

# ── Page 3, inventaire et selection ─────────────────────────────────────────
doc.add_page_break()
titre("Inventaire et sélection",
      "Laisser le client lister ses tâches sans rien lui suggérer. Estimer ensuite à la volée "
      "fois par an et minutes, puis heures par an. Les compléments de l'enquêteur viennent après, "
      "dans les lignes du bas, marqués d'une étoile.")
horloge()
grille(["Tâche, dans ses mots", "Qui", "Fois / an", "Min", "H / an", "Retenue"], 16, 0.75, [7.4, 2.6, 2, 1.6, 1.8, 2])
doc.add_paragraph()
table2([
    ("Total H / an de l'inventaire", ""),
    ("Heures déclarées par semaine × 52", ""),
    ("Écart", "au-delà de 25 %, reprendre l'inventaire avant de sélectionner"),
    ("Parcours", "tâches qui s'enchaînent sur la même occurrence, trois au moins, "
                 "à repérer et à nommer, elles se chiffrent en un seul parcours"),
    ("Règle de sélection", "Les 4 tâches ou parcours les plus lourds en heures par an, plus celle qui agace le plus "
                           "si elle n'y est pas. Cinq au plus."),
    ("Celle qui agace le plus", ""),
    ("Celle qui fait perdre de l'argent en retard", ""),
], 0.9)

# ── Une page par tache retenue ──────────────────────────────────────────────
for n in range(1, 6):
    doc.add_page_break()
    titre(f"Tâche {n}", "Partie B de la fiche. Douze minutes au plus.")
    horloge()
    table2([
        ("Nom, dans ses mots", ""),
        ("Déclencheur", "qu'est-ce qui la démarre"),
        ("Fois sur les 12 derniers mois", "un nombre, saisons comprises"),
        ("Fois sur les 12 prochains mois", "estimation du client, plafonnée au calcul à 3 fois le passé"),
        ("Ce volume est-il engagé", "non ☐   oui ☐, preuve, dates fixées, inscriptions, contrats signés"),
        ("Durée active en minutes", "la dernière fois"),
        ("Temps pour s'y remettre et s'en défaire", "20 % par défaut, corrigé si le client le dit   0 ☐  10 % ☐  30 % ☐"),
        ("Acteur et son taux", ""),
        ("Parcours", "aucun / lequel"),
        ("D'où viennent les informations", "outil, fichier, tête"),
        ("Ce que ça produit", ""),
        ("Qui le reçoit", ""),
        ("Chaque cas est-il différent", ""),
        ("Si elle disparaissait demain", "qui le remarquerait"),
    ])
    intertitre("Les quatre questions qui décident, partie C")
    table2([
        ("Règle ou jugement", "règle ☐      jugement ☐"),
        ("La règle est écrite quelque part", "oui ☐      non ☐"),
        ("Les données existent déjà", "oui ☐      non ☐"),
        ("La sortie garde la même forme", "oui ☐      non ☐"),
        ("Coût d'une erreur non détectée", "faible ☐      élevé ☐"),
    ], 0.9)

# ── Fin ─────────────────────────────────────────────────────────────────────
doc.add_page_break()
titre("Fin d'entretien")
horloge()
table2([
    ("Adresse pour le devis", ""),
    ("Ce qui a surpris", "note libre de l'enquêteur"),
    ("Phrase de clôture", "Rapport sous 72 heures, rien ne se promet en séance."),
])

sortie = dossier / "grille-entretien.docx"
doc.save(sortie)
print(f"ecrit {sortie}, methode v{VERSION}")
